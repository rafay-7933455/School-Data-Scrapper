"""
PESRP SIS Data Scraper - Streamlit Dashboard
============================================
Fetches school statistics (enrollment, teachers, school level) from the
Punjab Education Sector Reform Programme (PESRP) SIS portal using EMIS codes.

Deploy on Streamlit Community Cloud: app.py + requirements.txt in the repo root.
"""

import inspect
import io
import re
import time
from typing import Any, Dict, List, Optional

import pandas as pd
import requests
import streamlit as st
import urllib3
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ==========================================
# Configuration & Constants
# ==========================================
BASE_URL = "https://sis.pesrp.edu.pk"
DASHBOARD_URL = f"{BASE_URL}/dashboard"
ENROLLMENT_URL = f"{BASE_URL}/dashboard_revamp/get_gender_summary_pie"
TEACHER_URL = f"{BASE_URL}/dashboard_revamp/get_teacher_designation_stats"
SCHOOL_LEVEL_URL = f"{BASE_URL}/dashboard_revamp/get_school_levels_count"

DEFAULT_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json, text/javascript, */*; q=0.01",
    "X-Requested-With": "XMLHttpRequest",
}

DEFAULT_EMIS_CODES = [
    34510001,
    34510002,
    34510003,
    34510004,
    34510006,
    34510007,
    34510008,
    34510009,
    34510010,
    34510011,
    34510012,
    34510013,
    34510014,
    34510015,
    34510016,
    34510017,
    34510018,
    34510019,
    34510020,
    34510021,
    34510022,
    34510023,
    34510024,
    34510027,
    34510028,
    34510029,
    34510031,
    34510032,
    34510033,
    34510045,
    34510046,
    34510048,
    34510055,
    34510060,
    34510061,
    34510062,
    34510063,
    34510064,
    34510065,
    34510068,
    34510287,
    34510457,
    34510474,
    34510486,
    34510494,
    34510515,
    34510602,
    34510604,
    34510622,
    34510623,
    34510659,
    34520001,
    34520002,
    34520003,
    34520005,
    34520006,
    34520007,
    34520008,
    34520009,
    34520010,
    34520011,
    34520012,
    34520013,
    34520014,
    34520015,
    34520016,
    34520017,
    34520018,
    34520020,
    34520023,
    34520024,
    34520025,
    34520028,
    34520030,
    34520031,
    34520032,
    34520033,
    34520034,
    34520037,
    34520039,
    34520042,
    34520043,
    34520160,
    34520201,
    34520235,
    34520415,
    34520417,
    34520418,
    34520419,
    34520616,
    34520620,
]


# ==========================================
# Streamlit version compatibility
# ==========================================
def _stretch(func) -> Dict[str, Any]:
    """Full-width kwarg that works on both old and new Streamlit versions."""
    try:
        params = inspect.signature(func).parameters
    except (TypeError, ValueError):
        return {}
    if "width" in params:
        return {"width": "stretch"}
    if "use_container_width" in params:
        return {"use_container_width": True}
    return {}


BTN_FULL = _stretch(st.button)
DL_FULL = _stretch(st.download_button)
DF_FULL = _stretch(st.dataframe)


# ==========================================
# Helpers
# ==========================================
def parse_int_value(val: Any) -> int:
    """Safely convert string/numeric values into integers."""
    if val is None:
        return 0
    if isinstance(val, bool):
        return int(val)
    if isinstance(val, (int, float)):
        return int(val)
    if isinstance(val, str):
        cleaned = re.sub(r"[^\d-]", "", val.strip())
        try:
            return int(cleaned) if cleaned else 0
        except ValueError:
            return 0
    return 0


class PESRPScraper:
    """HTTP session management and scraping for the PESRP SIS portal."""

    def __init__(
        self,
        request_timeout: int = 20,
        verify_ssl: bool = True,
        retries: int = 2,
        proxy: Optional[str] = None,
    ):
        self.session = requests.Session()
        self.session.headers.update(DEFAULT_HEADERS)
        if proxy:
            self.session.proxies = {"http": proxy, "https": proxy}
        self.timeout = request_timeout
        self.verify = verify_ssl
        self.retries = retries
        self.csrf_token: Optional[str] = None

        if not verify_ssl:
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    # ---------- session / CSRF ----------
    def initialize_session(self) -> str:
        """Open the dashboard, store cookies and read the CSRF token."""
        response = self.session.get(
            DASHBOARD_URL,
            timeout=self.timeout,
            verify=self.verify,
            headers={"Accept": "text/html,application/xhtml+xml"},
        )
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")
        token = None

        csrf_input = soup.find("input", {"name": "csrf_test_name"})
        if csrf_input and csrf_input.get("value"):
            token = csrf_input["value"]
        else:
            csrf_meta = soup.find("meta", {"name": "csrf_test_name"})
            if csrf_meta and csrf_meta.get("content"):
                token = csrf_meta["content"]

        if not token:
            token = self._token_from_cookies()

        if not token:
            raise ValueError(
                "CSRF token not found. The portal may be blocking this server "
                "(e.g. non-Pakistan IPs) or its page layout has changed."
            )

        self.csrf_token = token
        return token

    def _token_from_cookies(self) -> Optional[str]:
        """CodeIgniter keeps the current CSRF token in a cookie too."""
        for cookie in self.session.cookies:
            if "csrf" in cookie.name.lower() and cookie.value:
                return cookie.value
        return None

    # ---------- requests ----------
    def _post_chart_endpoint(self, endpoint_url: str, emis_code: str) -> Dict[str, Any]:
        """POST to a chart endpoint with CSRF handling, retries and safe JSON parsing."""
        last_error: Optional[Exception] = None

        for attempt in range(self.retries + 1):
            try:
                if not self.csrf_token:
                    self.initialize_session()

                payload = {
                    "csrf_test_name": self.csrf_token,
                    "wing": "",
                    "district": "",
                    "tehsil": "",
                    "markaz": "",
                    "school": "",
                    "classes": "",
                    "s_id_emis_code": str(emis_code),
                }

                response = self.session.post(
                    endpoint_url,
                    data=payload,
                    headers={"Referer": DASHBOARD_URL, "Origin": BASE_URL},
                    timeout=self.timeout,
                    verify=self.verify,
                )

                # Token expired / rejected -> refresh session and retry
                if response.status_code in (403, 419):
                    self.csrf_token = None
                    raise requests.HTTPError(
                        f"HTTP {response.status_code} (CSRF rejected)"
                    )

                response.raise_for_status()

                try:
                    data = response.json()
                except ValueError:
                    # HTML instead of JSON usually means expired session
                    self.csrf_token = None
                    raise ValueError("Server returned non-JSON response")

                if not isinstance(data, dict):
                    data = {"data": data}

                # Token rotation: prefer the one in the JSON, else the cookie
                new_token = (
                    data.get("csrf_test_name")
                    or data.get("csrf_token")
                    or self._token_from_cookies()
                )
                if new_token:
                    self.csrf_token = new_token

                return data

            except Exception as exc:  # noqa: BLE001
                last_error = exc
                time.sleep(0.5 * (attempt + 1))

        raise RuntimeError(str(last_error))

    # ---------- public API ----------
    def fetch_school_record(
        self, emis_code: Any, fetch_extended: bool = True
    ) -> Dict[str, Any]:
        """Fetch all stats for one EMIS code."""
        emis_str = str(emis_code).strip()

        enrollment_resp = self._post_chart_endpoint(ENROLLMENT_URL, emis_str)

        boys = parse_int_value(
            enrollment_resp.get("male") or enrollment_resp.get("boys")
        )
        girls = parse_int_value(
            enrollment_resp.get("female") or enrollment_resp.get("girls")
        )
        total = parse_int_value(enrollment_resp.get("total")) or (boys + girls)

        record: Dict[str, Any] = {"EMIS Code": emis_str, "Total Enrollment": total}

        if fetch_extended:
            record["Boys"] = boys
            record["Girls"] = girls

            try:
                teacher_resp = self._post_chart_endpoint(TEACHER_URL, emis_str)
                record["Total Teachers"] = parse_int_value(teacher_resp.get("total"))
            except Exception:  # noqa: BLE001
                record["Total Teachers"] = None

            try:
                level_resp = self._post_chart_endpoint(SCHOOL_LEVEL_URL, emis_str)
                record["School Level"] = (
                    level_resp.get("level_name")
                    or level_resp.get("school_level")
                    or "N/A"
                )
            except Exception:  # noqa: BLE001
                record["School Level"] = "N/A"

        return record


# ==========================================
# Excel export
# ==========================================
def export_to_styled_excel(df: pd.DataFrame) -> bytes:
    """Generate a styled Excel file in memory."""
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="School Stats")
        worksheet = writer.sheets["School Stats"]

        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        row_alt_fill = PatternFill(
            start_color="F2F5F9", end_color="F2F5F9", fill_type="solid"
        )
        side = Side(style="thin", color="D9D9D9")
        border_thin = Border(left=side, right=side, top=side, bottom=side)
        center_align = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row_idx, row in enumerate(
            worksheet.iter_rows(
                min_row=2, max_row=len(df) + 1, min_col=1, max_col=len(df.columns)
            ),
            start=2,
        ):
            for cell in row:
                cell.border = border_thin
                if row_idx % 2 == 0:
                    cell.fill = row_alt_fill
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
                else:
                    cell.alignment = center_align

        for col in worksheet.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=0)
            worksheet.column_dimensions[get_column_letter(col[0].column)].width = max(
                max_len + 4, 12
            )

        worksheet.freeze_panes = "A2"

    return output.getvalue()


# ==========================================
# Streamlit UI
# ==========================================
def main() -> None:
    st.set_page_config(
        page_title="PESRP School Data Scraper",
        page_icon="🏫",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    st.title("🏫 PESRP SIS School Data Extractor")
    st.markdown(
        "Extract enrollment, teacher stats, and school levels from the "
        "**Punjab Education Sector Reform Programme (PESRP)** portal."
    )

    # ---------------- Sidebar ----------------
    st.sidebar.header("⚙️ Configuration")

    input_mode = st.sidebar.radio(
        "EMIS Input Source:",
        [f"Use Default List ({len(DEFAULT_EMIS_CODES)} Schools)", "Custom EMIS Codes"],
    )

    if input_mode == "Custom EMIS Codes":
        raw_input = st.sidebar.text_area(
            "Enter EMIS Codes (comma, space, or new line separated):",
            value="34510001\n34510002\n34510003",
            height=150,
        )
        emis_list: List[int] = list(
            dict.fromkeys(int(c) for c in re.findall(r"\d+", raw_input))
        )
    else:
        emis_list = DEFAULT_EMIS_CODES

    st.sidebar.subheader("Extraction Options")
    fetch_extended = st.sidebar.checkbox(
        "Fetch Extended Data (Teachers, Gender, Level)",
        value=True,
        help="Retrieves extra stats (3x more requests per school).",
    )
    request_delay = st.sidebar.slider(
        "Request Delay (seconds):",
        0.0,
        2.0,
        0.3,
        0.1,
        help="Delay between schools to avoid rate-limiting.",
    )
    verify_ssl = st.sidebar.checkbox(
        "Verify SSL certificate",
        value=True,
        help="Untick only if you get SSL certificate errors from the portal.",
    )

    try:
        secret_proxy = st.secrets.get("PROXY_URL", "")
    except Exception:  # noqa: BLE001
        secret_proxy = ""
    proxy_url = st.sidebar.text_input(
        "Proxy URL (optional):",
        value=secret_proxy,
        type="password",
        placeholder="http://user:pass@host:port",
        help="Use a Pakistan-based proxy if the portal blocks Streamlit Cloud's servers.",
    ).strip()

    st.sidebar.markdown("---")
    st.sidebar.info(f"📍 **Total EMIS codes to process:** `{len(emis_list)}`")

    # ---------------- Run ----------------
    start_btn = st.button("🚀 Start Scraping Data", type="primary", **BTN_FULL)

    if "scraped_df" not in st.session_state:
        st.session_state.scraped_df = None
    if "failed" not in st.session_state:
        st.session_state.failed = []

    if start_btn:
        if not emis_list:
            st.error("Please provide at least one valid numeric EMIS code.")
            return

        scraper = PESRPScraper(verify_ssl=verify_ssl, proxy=proxy_url or None)
        scraped_records: List[Dict[str, Any]] = []
        failed: List[Dict[str, str]] = []

        with st.status("Connecting to PESRP SIS portal...", expanded=True) as status:
            try:
                scraper.initialize_session()
                st.write("✅ Session established and CSRF token retrieved.")
            except Exception as e:  # noqa: BLE001
                status.update(label="❌ Failed to connect to server", state="error")
                if isinstance(e, requests.exceptions.SSLError):
                    hint = "SSL certificate problem. Untick 'Verify SSL certificate' in the sidebar."
                elif isinstance(e, requests.exceptions.ConnectTimeout):
                    hint = "Connection timed out. Typical of a firewall silently dropping foreign IPs."
                elif isinstance(e, requests.exceptions.ReadTimeout):
                    hint = "Server accepted the connection but never answered (slow or blocking)."
                elif isinstance(e, requests.exceptions.ConnectionError):
                    hint = "Connection refused or reset. Typical of IP/geo blocking."
                elif isinstance(e, requests.exceptions.HTTPError):
                    hint = (
                        "Server responded with an HTTP error (e.g. 403 means blocked)."
                    )
                else:
                    hint = "Unexpected response from the portal."
                st.error(f"{type(e).__name__}: {e}")
                st.info(
                    f"**Diagnosis:** {hint}\n\n"
                    "If this works on your own computer but not on Streamlit Cloud, "
                    "the portal is blocking the cloud server's non-Pakistan IP. "
                    "Use a Pakistan-based proxy (sidebar) or run the app locally."
                )
                return

            progress_bar = st.progress(0.0)
            status_text = st.empty()
            total_count = len(emis_list)

            for idx, emis in enumerate(emis_list, start=1):
                status_text.text(f"Fetching [{idx}/{total_count}]: EMIS {emis}")
                try:
                    scraped_records.append(
                        scraper.fetch_school_record(emis, fetch_extended=fetch_extended)
                    )
                except Exception as err:  # noqa: BLE001
                    failed.append({"EMIS Code": str(emis), "Error": str(err)})

                progress_bar.progress(idx / total_count)
                if request_delay > 0:
                    time.sleep(request_delay)

            label = f"🎉 Scraped {len(scraped_records)} of {total_count} schools"
            status.update(
                label=label,
                state="complete" if scraped_records else "error",
            )

        st.session_state.failed = failed
        if scraped_records:
            df = (
                pd.DataFrame(scraped_records)
                .sort_values("EMIS Code")
                .reset_index(drop=True)
            )
            st.session_state.scraped_df = df
        else:
            st.session_state.scraped_df = None
            st.error("No records could be fetched. See errors below.")

    # ---------------- Results ----------------
    df = st.session_state.scraped_df
    if df is not None:
        st.markdown("---")
        st.subheader("📊 Scraped Results")

        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Total Schools Scraped", len(df))
        m2.metric("Total Enrollment", f"{int(df['Total Enrollment'].sum()):,}")
        if "Boys" in df.columns and "Girls" in df.columns:
            m3.metric("Total Boys", f"{int(df['Boys'].sum()):,}")
            m4.metric("Total Girls", f"{int(df['Girls'].sum()):,}")

        st.dataframe(df, hide_index=True, **DF_FULL)

        col_dl1, col_dl2 = st.columns(2)
        col_dl1.download_button(
            label="📥 Download Excel Report (.xlsx)",
            data=export_to_styled_excel(df),
            file_name="PESRP_School_Statistics.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            **DL_FULL,
        )
        col_dl2.download_button(
            label="📄 Download CSV File (.csv)",
            data=df.to_csv(index=False).encode("utf-8"),
            file_name="PESRP_School_Statistics.csv",
            mime="text/csv",
            **DL_FULL,
        )

    if st.session_state.failed:
        with st.expander(f"⚠️ {len(st.session_state.failed)} EMIS codes failed"):
            st.dataframe(
                pd.DataFrame(st.session_state.failed), hide_index=True, **DF_FULL
            )


if __name__ == "__main__":
    main()
